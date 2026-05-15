#!/usr/bin/env python3
"""
作図チェック工程表生成ツール
施工工程表を読み込み、施工図作成〜設計事務所確認までの作図チェック工程表をExcelで出力する。
"""

import argparse
import sys
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side
)
from openpyxl.utils import get_column_letter


# ──────────────────────────────────────────────
# データ定義
# ──────────────────────────────────────────────

@dataclass
class WorkItem:
    """施工工程の1件"""
    name: str           # 工種名
    area: str           # 工区・部位
    start: date         # 施工開始日
    end: date           # 施工完了日


@dataclass
class PhaseConfig:
    """作図チェック各フェーズの設定"""
    name: str
    days: int           # 期間（営業日ではなく暦日）
    color: str          # セル塗りつぶし色 (ARGB hex)


# デフォルトフェーズ設定（施工開始から逆算）
DEFAULT_PHASES = [
    PhaseConfig("施工図作成",          days=14, color="FF92D050"),   # 緑
    PhaseConfig("社内チェック",        days=7,  color="FFFFC000"),   # 黄
    PhaseConfig("設計事務所提出・確認", days=10, color="FF00B0F0"),   # 青
]

# 施工開始の何日前までに全フェーズを終わらせるか（バッファ）
BUFFER_DAYS_BEFORE_START = 3


# ──────────────────────────────────────────────
# 入力
# ──────────────────────────────────────────────

def load_from_excel(path: Path) -> list[WorkItem]:
    """
    Excelから施工工程を読み込む。
    想定フォーマット（1行目ヘッダー）:
      A: 工種名, B: 工区/部位, C: 施工開始日, D: 施工完了日
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name, area, start, end = row[0], row[1], row[2], row[3]
        if not name or not start or not end:
            continue
        if isinstance(start, str):
            start = date.fromisoformat(start)
        if isinstance(end, str):
            end = date.fromisoformat(end)
        if hasattr(start, "date"):
            start = start.date()
        if hasattr(end, "date"):
            end = end.date()
        items.append(WorkItem(name=str(name), area=str(area or ""), start=start, end=end))
    return items


def sample_data() -> list[WorkItem]:
    """デモ用サンプル施工工程データ"""
    base = date(2025, 9, 1)
    return [
        WorkItem("躯体工事（1F）",    "1階",    base,                  base + timedelta(30)),
        WorkItem("躯体工事（2F）",    "2階",    base + timedelta(25),  base + timedelta(60)),
        WorkItem("外装工事",          "外部",   base + timedelta(50),  base + timedelta(90)),
        WorkItem("内装工事（1F）",    "1階",    base + timedelta(45),  base + timedelta(80)),
        WorkItem("内装工事（2F）",    "2階",    base + timedelta(60),  base + timedelta(95)),
        WorkItem("設備工事（給排水）", "全体",   base + timedelta(30),  base + timedelta(85)),
        WorkItem("設備工事（電気）",   "全体",   base + timedelta(35),  base + timedelta(90)),
        WorkItem("外構工事",          "外部",   base + timedelta(80),  base + timedelta(110)),
    ]


# ──────────────────────────────────────────────
# スケジュール計算
# ──────────────────────────────────────────────

@dataclass
class PhaseSchedule:
    phase: PhaseConfig
    start: date
    end: date


@dataclass
class DrawingCheckRow:
    work: WorkItem
    phases: list[PhaseSchedule]

    @property
    def check_start(self) -> date:
        return self.phases[0].start

    @property
    def check_end(self) -> date:
        return self.phases[-1].end


def calc_drawing_schedule(
    items: list[WorkItem],
    phases: list[PhaseConfig],
    buffer: int = BUFFER_DAYS_BEFORE_START,
) -> list[DrawingCheckRow]:
    """各工種について作図チェック工程を逆算して返す"""
    rows = []
    for item in items:
        deadline = item.start - timedelta(buffer)
        # 末尾フェーズから逆順に日付を割り当てる
        phase_schedules = []
        current_end = deadline
        for phase in reversed(phases):
            p_end = current_end
            p_start = p_end - timedelta(phase.days - 1)
            phase_schedules.insert(0, PhaseSchedule(phase, p_start, p_end))
            current_end = p_start - timedelta(1)
        rows.append(DrawingCheckRow(work=item, phases=phase_schedules))
    return rows


# ──────────────────────────────────────────────
# Excel出力
# ──────────────────────────────────────────────

HEADER_FILL   = PatternFill("solid", fgColor="FF1F3864")
HEADER_FONT   = Font(color="FFFFFFFF", bold=True, size=10)
SUBHDR_FILL   = PatternFill("solid", fgColor="FF2E75B6")
SUBHDR_FONT   = Font(color="FFFFFFFF", bold=True, size=9)
LABEL_FONT    = Font(size=9)
WEEKEND_FILL  = PatternFill("solid", fgColor="FFF2F2F2")
THIN_BORDER   = Border(
    left=Side(style="thin", color="FFD9D9D9"),
    right=Side(style="thin", color="FFD9D9D9"),
    top=Side(style="thin", color="FFD9D9D9"),
    bottom=Side(style="thin", color="FFD9D9D9"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

# 固定列の列番号（1始まり）
COL_NO      = 1
COL_AREA    = 2
COL_WORK    = 3
COL_CONST_S = 4   # 施工開始
COL_CONST_E = 5   # 施工完了
COL_DATA_START = 6  # ガント開始列


def _make_date_range(rows: list[DrawingCheckRow]) -> list[date]:
    """表示するカレンダー範囲を生成"""
    all_dates = []
    for row in rows:
        all_dates.append(row.check_start)
        all_dates.append(row.work.end)
    min_d = min(all_dates)
    max_d = max(all_dates)
    # 月の始まりに揃える
    min_d = min_d.replace(day=1)
    # 月末まで延ばす
    if max_d.month == 12:
        max_d = max_d.replace(year=max_d.year + 1, month=1, day=1) - timedelta(1)
    else:
        max_d = max_d.replace(month=max_d.month + 1, day=1) - timedelta(1)

    days = []
    d = min_d
    while d <= max_d:
        days.append(d)
        d += timedelta(1)
    return days


def _col_for_date(date_range: list[date], d: date) -> int | None:
    """日付 → 列インデックス（0始まり）を返す。範囲外は None"""
    if d < date_range[0] or d > date_range[-1]:
        return None
    return (d - date_range[0]).days


def write_excel(
    rows: list[DrawingCheckRow],
    phases: list[PhaseConfig],
    output_path: Path,
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "作図チェック工程表"

    date_range = _make_date_range(rows)
    total_days = len(date_range)

    # ── 列幅設定 ──────────────────────────────
    ws.column_dimensions[get_column_letter(COL_NO)].width      = 4
    ws.column_dimensions[get_column_letter(COL_AREA)].width     = 10
    ws.column_dimensions[get_column_letter(COL_WORK)].width     = 18
    ws.column_dimensions[get_column_letter(COL_CONST_S)].width  = 11
    ws.column_dimensions[get_column_letter(COL_CONST_E)].width  = 11
    for i in range(total_days):
        ws.column_dimensions[get_column_letter(COL_DATA_START + i)].width = 2.2

    # ── 行1: タイトル ─────────────────────────
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1, end_column=COL_DATA_START + total_days - 1)
    title_cell = ws.cell(1, 1, "作図チェック工程表")
    title_cell.font = Font(bold=True, size=14, color="FF1F3864")
    title_cell.alignment = CENTER

    # ── 行2: 月ヘッダー ──────────────────────
    ws.row_dimensions[2].height = 16
    month_start_col = COL_DATA_START
    prev_month = None
    for i, d in enumerate(date_range):
        if d.month != prev_month:
            if prev_month is not None:
                # 前の月をマージ
                end_col = COL_DATA_START + i - 1
                if end_col > month_start_col:
                    ws.merge_cells(start_row=2, start_column=month_start_col,
                                   end_row=2, end_column=end_col)
                ws.cell(2, month_start_col).alignment = CENTER
            month_start_col = COL_DATA_START + i
            prev_month = d.month
    # 最後の月
    end_col = COL_DATA_START + total_days - 1
    if end_col > month_start_col:
        ws.merge_cells(start_row=2, start_column=month_start_col,
                       end_row=2, end_column=end_col)

    # 月ヘッダー文字の記入
    for i, d in enumerate(date_range):
        if d.day == 1:
            c = ws.cell(2, COL_DATA_START + i, f"{d.year}年{d.month}月")
            c.fill = SUBHDR_FILL
            c.font = SUBHDR_FONT
            c.alignment = CENTER

    # ── 行3: 日付ヘッダー ────────────────────
    ws.row_dimensions[3].height = 14
    for i, d in enumerate(date_range):
        c = ws.cell(3, COL_DATA_START + i, d.day)
        c.font = Font(size=7, bold=(d.weekday() >= 5))
        c.alignment = CENTER
        c.border = THIN_BORDER
        if d.weekday() >= 5:
            c.fill = WEEKEND_FILL
            c.font = Font(size=7, bold=True,
                          color="FFFF0000" if d.weekday() == 6 else "FF0070C0")

    # ── 行4: 曜日ヘッダー ────────────────────
    ws.row_dimensions[4].height = 12
    DOW = ["月", "火", "水", "木", "金", "土", "日"]
    for i, d in enumerate(date_range):
        c = ws.cell(4, COL_DATA_START + i, DOW[d.weekday()])
        c.font = Font(size=7, bold=(d.weekday() >= 5),
                      color="FFFF0000" if d.weekday() == 6
                      else ("FF0070C0" if d.weekday() == 5 else "FF000000"))
        c.alignment = CENTER
        c.border = THIN_BORDER
        if d.weekday() >= 5:
            c.fill = WEEKEND_FILL

    # ── 行5: 固定列ヘッダー ──────────────────
    ws.row_dimensions[5].height = 20
    headers = ["No.", "工区/部位", "工種名", "施工開始", "施工完了"]
    for col_idx, label in enumerate(headers, start=1):
        c = ws.cell(5, col_idx, label)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = THIN_BORDER

    # 凡例（固定列ヘッダー行にフェーズ色凡例を右端に追加）
    legend_col = COL_DATA_START + total_days + 1
    ws.cell(5, legend_col, "凡例").font = Font(bold=True, size=9)
    for pi, ph in enumerate(phases):
        lc = ws.cell(5, legend_col + 1 + pi, ph.name)
        lc.fill = PatternFill("solid", fgColor=ph.color)
        lc.font = Font(size=8, bold=True)
        lc.alignment = CENTER
        ws.column_dimensions[get_column_letter(legend_col + 1 + pi)].width = 16

    # ── 週末の列を薄くする ───────────────────
    for i, d in enumerate(date_range):
        if d.weekday() >= 5:
            for row_idx in range(6, 6 + len(rows)):
                c = ws.cell(row_idx, COL_DATA_START + i)
                c.fill = WEEKEND_FILL

    # ── データ行 ─────────────────────────────
    for row_idx, row in enumerate(rows, start=6):
        ws.row_dimensions[row_idx].height = 18

        # 固定列
        def fixed(col, val, align=CENTER):
            c = ws.cell(row_idx, col, val)
            c.font = LABEL_FONT
            c.alignment = align
            c.border = THIN_BORDER
            return c

        fixed(COL_NO,      row_idx - 5)
        fixed(COL_AREA,    row.work.area,  LEFT)
        fixed(COL_WORK,    row.work.name,  LEFT)
        fixed(COL_CONST_S, row.work.start.strftime("%Y/%m/%d"))
        fixed(COL_CONST_E, row.work.end.strftime("%Y/%m/%d"))

        # ガントバー：作図フェーズ
        for ps in row.phases:
            si = _col_for_date(date_range, ps.start)
            ei = _col_for_date(date_range, ps.end)
            if si is None or ei is None:
                continue
            fill = PatternFill("solid", fgColor=ps.phase.color)
            for di in range(si, ei + 1):
                c = ws.cell(row_idx, COL_DATA_START + di)
                c.fill = fill
                c.border = THIN_BORDER

        # ガントバー：施工期間（薄い灰色のハッチ）
        si = _col_for_date(date_range, row.work.start)
        ei = _col_for_date(date_range, row.work.end)
        if si is not None and ei is not None:
            for di in range(si, ei + 1):
                c = ws.cell(row_idx, COL_DATA_START + di)
                c.fill = PatternFill("solid", fgColor="FFD6DCE4")
                c.border = THIN_BORDER

    # ── 凡例（施工期間）────────────────────
    sc = ws.cell(5, legend_col + 1 + len(phases), "施工期間")
    sc.fill = PatternFill("solid", fgColor="FFD6DCE4")
    sc.font = Font(size=8, bold=True)
    sc.alignment = CENTER
    ws.column_dimensions[get_column_letter(legend_col + 1 + len(phases))].width = 10

    wb.save(output_path)
    print(f"✅ 出力完了: {output_path}")


# ──────────────────────────────────────────────
# CLI
# ──────────────────────────────────────────────

def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="施工工程表から作図チェック工程表を生成します"
    )
    p.add_argument(
        "-i", "--input",
        help="入力Excelファイル（省略時はサンプルデータを使用）",
        type=Path,
        default=None,
    )
    p.add_argument(
        "-o", "--output",
        help="出力Excelファイル（デフォルト: drawing_check_schedule.xlsx）",
        type=Path,
        default=Path("drawing_check_schedule.xlsx"),
    )
    p.add_argument(
        "--drawing-days",
        help="施工図作成期間（日数、デフォルト: 14）",
        type=int, default=14,
    )
    p.add_argument(
        "--check-days",
        help="社内チェック期間（日数、デフォルト: 7）",
        type=int, default=7,
    )
    p.add_argument(
        "--review-days",
        help="設計事務所提出・確認期間（日数、デフォルト: 10）",
        type=int, default=10,
    )
    p.add_argument(
        "--buffer-days",
        help="施工開始前バッファ日数（デフォルト: 3）",
        type=int, default=BUFFER_DAYS_BEFORE_START,
    )
    return p.parse_args()


def main() -> None:
    args = parse_args()

    phases = [
        PhaseConfig("施工図作成",           args.drawing_days, "FF92D050"),
        PhaseConfig("社内チェック",         args.check_days,   "FFFFC000"),
        PhaseConfig("設計事務所提出・確認", args.review_days,  "FF00B0F0"),
    ]

    if args.input:
        if not args.input.exists():
            print(f"❌ ファイルが見つかりません: {args.input}", file=sys.stderr)
            sys.exit(1)
        print(f"📂 工程表を読み込み中: {args.input}")
        items = load_from_excel(args.input)
        if not items:
            print("❌ データが読み込めませんでした。ヘッダー行を除いたデータが存在するか確認してください。",
                  file=sys.stderr)
            sys.exit(1)
    else:
        print("📋 サンプルデータを使用します")
        items = sample_data()

    print(f"   {len(items)} 件の工種を処理します")

    rows = calc_drawing_schedule(items, phases, args.buffer_days)
    write_excel(rows, phases, args.output)

    print()
    print("【フェーズ設定】")
    for ph in phases:
        print(f"  {ph.name}: {ph.days}日間")
    print(f"  施工開始前バッファ: {args.buffer_days}日")


if __name__ == "__main__":
    main()
